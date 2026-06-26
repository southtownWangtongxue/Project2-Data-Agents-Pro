"""
文件导出接口
支持将查询结果导出为 Excel 或 CSV 格式
"""
from fastapi import APIRouter
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field
import io
import csv
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

router = APIRouter(prefix="/export", tags=["export"])


class ExportRequest(BaseModel):
    """导出请求体"""
    data: list[dict] = Field(..., description="要导出的数据行列表")
    columns: list[str] = Field(..., description="列名列表")
    format: str = Field(default="csv", description="导出格式: csv 或 excel")
    filename: str = Field(default="export", description="文件名（不含扩展名）")


@router.post("/csv", summary="导出为 CSV 格式")
async def export_csv(body: ExportRequest):
    """
    将查询结果导出为 CSV 文件并返回文件流。
    """
    output = io.StringIO()
    if body.columns:
        writer = csv.DictWriter(output, fieldnames=body.columns)
        writer.writeheader()
        writer.writerows(body.data)
    else:
        writer = csv.writer(output)
        for row in body.data:
            writer.writerow(row.values())

    output.seek(0)

    return StreamingResponse(
        iter([output.getvalue()]),
        media_type="text/csv",
        headers={
            "Content-Disposition": f"attachment; filename={body.filename}.csv",
        },
    )


@router.post("/excel", summary="导出为 Excel 格式")
async def export_excel(body: ExportRequest):
    """
    将查询结果导出为带样式的 .xlsx 文件并返回文件流。

    功能：
    - 表头深色背景 + 白色粗体文字 + 自动列宽
    - 数据行交替颜色（斑马条纹）
    - 数字列自动右对齐
    - 底部汇总行（行数统计）
    """
    wb = Workbook()
    ws = wb.active
    ws.title = body.filename[:31] or "Sheet1"

    if not body.columns:
        # 无边界的快速路径
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return StreamingResponse(
            iter([output.getvalue()]),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={
                "Content-Disposition": f"attachment; filename={body.filename}.xlsx",
            },
        )

    # ── 样式定义 ──────────────────────────────────
    header_fill = PatternFill(start_color="6366F1", end_color="6366F1", fill_type="solid")
    header_font = Font(name="Microsoft YaHei", bold=True, color="FFFFFF", size=11)
    header_align = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style="thin", color="D1D5DB"),
        right=Side(style="thin", color="D1D5DB"),
        top=Side(style="thin", color="D1D5DB"),
        bottom=Side(style="thin", color="D1D5DB"),
    )
    even_fill = PatternFill(start_color="F3F4F6", end_color="F3F4F6", fill_type="solid")
    data_font = Font(name="Microsoft YaHei", size=10)
    data_align_left = Alignment(horizontal="left", vertical="center")
    data_align_right = Alignment(horizontal="right", vertical="center")

    # ── 写表头 ──────────────────────────────────
    for col_idx, col_name in enumerate(body.columns, 1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_align
        cell.border = thin_border

    # ── 判定各列是否为数字列（取前 100 行采样） ──
    numeric_cols: set[int] = set()
    for row_data in body.data[:100]:
        for col_idx, col_name in enumerate(body.columns, 1):
            val = row_data.get(col_name)
            if isinstance(val, (int, float)) and not isinstance(val, bool):
                numeric_cols.add(col_idx)

    # ── 写数据行 ─────────────────────────────────
    for row_idx, row_data in enumerate(body.data, 2):
        is_even = row_idx % 2 == 0
        for col_idx, col_name in enumerate(body.columns, 1):
            val = row_data.get(col_name)
            # 处理 datetime 类型
            if isinstance(val, datetime):
                val = val.strftime("%Y-%m-%d %H:%M:%S")
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = data_font
            cell.border = thin_border
            cell.alignment = data_align_right if col_idx in numeric_cols else data_align_left
            if is_even:
                cell.fill = even_fill

    # ── 底部汇总行 ───────────────────────────────
    summary_row = len(body.data) + 2
    summary_fill = PatternFill(start_color="EEF2FF", end_color="EEF2FF", fill_type="solid")
    summary_cell = ws.cell(row=summary_row, column=1, value=f"共 {len(body.data)} 条记录")
    summary_cell.font = Font(name="Microsoft YaHei", bold=True, size=10, color="6366F1")
    summary_cell.fill = summary_fill
    summary_cell.border = thin_border

    # ── 自动列宽（中文计 2，ASCII 计 1） ─────────
    for col_idx, col_name in enumerate(body.columns, 1):
        # 表头宽度
        header_width = sum(2 if ord(c) > 127 else 1 for c in str(col_name))
        max_width = header_width
        # 抽样数据宽度
        for row_data in body.data[:50]:
            val = str(row_data.get(col_name, ""))[:80]
            val_width = sum(2 if ord(c) > 127 else 1 for c in val)
            max_width = max(max_width, val_width)
        # 限制最小/最大宽度
        ws.column_dimensions[get_column_letter(col_idx)].width = max(10, min(max_width + 4, 40))

    # ── 冻结首行 ─────────────────────────────────
    ws.freeze_panes = "A2"

    # ── 写入流 ───────────────────────────────────
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return StreamingResponse(
        iter([output.getvalue()]),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f"attachment; filename={body.filename}.xlsx",
        },
    )
